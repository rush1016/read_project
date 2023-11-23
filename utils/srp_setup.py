from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

class SrpSetup():
    def set_grade(table):
        GRADE_LEVELS = ('IV', 'V', 'VI')
        row_index = 2
        col_index = 0
        interval = 7

        for grade_level in GRADE_LEVELS:
            cell = table.cell(row_index, col_index)
            cell.text = grade_level
            row_index += interval

    def set_sections(table, sections):
        section_col = 1
        CLASS_SECTIONS = {
            4: [
                f'{section.section_name}' for section in sections.filter(grade_level=4)
            ],
            5: [
                f'{section.section_name}' for section in sections.filter(grade_level=5)
            ],
            6: [
                f'{section.section_name}' for section in sections.filter(grade_level=6)
            ],
        }
        for grade, section_list in CLASS_SECTIONS.items():
            if grade == 4:
                start_row = 3

            elif grade == 5:
                start_row = 10

            elif grade == 6:
                start_row = 17

            current_row = start_row
            for section in section_list:
                table.cell(current_row, section_col).text = section
                current_row += 1


    def set_total_students_per_grade(table, students):
        grades = [4, 5, 6]
        start_row = 2

        totals_col = 2
        failed_total_col = 3
        passed_total_col = 4

        for i, grade in enumerate(grades, start=1):
            total_students = students.filter(grade_level=grade).count()
            total_failed = students.filter(grade_level=grade, gst_score__lt=14).count()
            total_passed = students.filter(grade_level=grade, gst_score__gte=14).count()

            row = start_row + (i - 1) * 7
            table.cell(row, totals_col).text = str(total_students)
            table.cell(row, failed_total_col).text = str(total_failed)
            table.cell(row, passed_total_col).text = str(total_passed)


    def set_total_per_section(table, students, sections):
        start_row = 3
        current_row = start_row
        totals_col = 2
        failed_col = 3
        passed_col = 4
        section_col = 1
        grade_row_spacing = 2


        for grade_level in [4, 5, 6]:
            total_per_section = {section.section_name: students.filter(class_section=section.section_name).count() for section in sections.filter(grade_level=grade_level)}
            failed_per_section = {section.section_name: students.filter(class_section=section.section_name, gst_score__lt=14).count() for section in sections.filter(grade_level=grade_level)}
            passed_per_section = {section.section_name: students.filter(class_section=section.section_name, gst_score__gte=14).count() for section in sections.filter(grade_level=grade_level)}

            for section, total in total_per_section.items():
                try:
                    section_cell = table.cell(current_row, section_col)
                    current_cell = table.cell(current_row, totals_col)
                    failed_cell = table.cell(current_row, failed_col)
                    passed_cell = table.cell(current_row, passed_col)

                    if section == section_cell.text:
                        current_cell.text = str(total)
                        failed_cell.text = str(failed_per_section[section])
                        passed_cell.text = str(passed_per_section[section])

                except IndexError:
                    # Handle the case where the cell indices are out of bounds
                    pass

                current_row += 1

            current_row += grade_row_spacing

    def set_school_total(table, students):
        footer1 = table.rows[-2]
        footer2 = table.rows[-1]

        total_col = 2
        failed_col = 3
        passed_col = 4

        total_students = str(students.all().count())
        total_failed = str(students.filter(gst_score__lt=14).count())
        total_passed = str(students.filter(gst_score__gte=14).count())

        footer1.cells[0].text = 'Total'
        paragraph = footer1.cells[0].paragraphs[0]
        run = paragraph.runs[0]
        run.bold = True

        footer2.cells[total_col].text = total_students
        footer2.cells[failed_col].text = total_failed
        footer2.cells[passed_col].text = total_passed